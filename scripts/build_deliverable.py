"""Build the fictitious deliverable offered for download in the demo.

    python scripts/build_deliverable.py

Writes, from src/data/catalogue.json:
    public/downloads/HarbourMarket_Protein_Ratio_DEMO.xlsx   formula-driven workbook
    public/downloads/HarbourMarket_products_DEMO.csv         line-level file

The workbook follows the layout of a real Altera deliverable: Welcome (with the volumes,
the only input), WWF, Protein Tracker, By family, By department, By brand, Top and flop,
and the line-level Data tab every figure is computed from. Everything in it is invented.

openpyxl writes formulas without evaluating them. Every formula cell is therefore
registered with the value Python expects; when LibreOffice is installed the workbook is
recalculated and each of those cells compared. The script refuses to finish on any difference.
"""
from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src" / "data" / "catalogue.json"
ASSETS = ROOT / "scripts" / "assets"
OUT = ROOT / "public" / "downloads"
XLSX = OUT / "HarbourMarket_Protein_Ratio_DEMO.xlsx"
CSV = OUT / "HarbourMarket_products_DEMO.csv"

# ------------------------------------------------------------ house style
NAVY, GREEN, ORANGE, GREY = "1F3864", "375623", "C55A11", "7F7F7F"
CLASS_FILL = {"animal_core": "C00000", "composite": "D99694", "plant_non_core": "A9D3AB", "plant_core": "375623"}
ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
NOTE = Font(name="Arial", size=10, color="404040")
SMALL = Font(name="Arial", size=9, italic=True, color=GREY)
LABEL = Font(name="Arial", size=9, bold=True, color=GREY)
NAVY_B = Font(name="Arial", size=11, bold=True, color=NAVY)
WHITE_B = Font(name="Arial", size=11, bold=True, color="FFFFFF")
KPI = Font(name="Arial", size=28, bold=True, color=GREEN)
WARN = Font(name="Arial", size=10, bold=True, color="C00000")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
BAND_FILL = PatternFill("solid", fgColor="EDF2F9")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_BORDER = Border(bottom=Side(style="medium", color=NAVY))
INPUT_BORDER = Border(bottom=Side(style="thin", color="BF8F00"))
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
T, PCT, INT, KG = "#,##0.0", "0.0%", "#,##0", "#,##0.0"
BANNER = "FICTITIOUS DATA — demo file. Harbour Market, its brands, its products and every figure are invented."

PT_ORDER = ["animal_core", "composite", "plant_non_core", "plant_core"]
PT_LABEL = {"animal_core": "Animal core", "composite": "Animal & plant composite",
            "plant_non_core": "Plant-based non-core", "plant_core": "Plant-based core"}
PT_SPLIT_LABEL = {"plant_core": "Plant-based core — plant, a protein source",
                  "plant_non_core": "Plant-based non-core — plant, not a protein source",
                  "composite": "Composite — plant and animal mixed", "animal_core": "Animal core — animal"}
# the level of detail the GPA methodology describes, class by class
DETAIL = {
    "animal_core": ["Red meat", "Poultry", "Fish and seafood", "Eggs", "Other animals (e.g. snails)", "Cheese",
                    "Milk, yoghurt and cream", "Dishes and processed products", "Discretionary products", "Other"],
    "composite": ["Meat-based dishes", "Fish-based dishes", "Vegetarian dishes", "Vegan dishes",
                  "Discretionary products", "Cereals and bread", "Other"],
    "plant_non_core": ["Refined grains", "Whole grains", "Fruit and vegetables", "Starchy foods",
                       "Discretionary products", "Outside the WWF scope", "Other"],
    "plant_core": ["Pulses", "Nuts", "Seeds", "Plant-based meat alternatives", "Tofu, tempeh, seitan",
                   "Plant-based dairy alternatives", "Other"],
}
SUB_LABEL = {"red_meat": "Red meat", "poultry": "Poultry", "fish_shellfish": "Fish & shellfish", "eggs": "Eggs",
             "other_animal": "Other animal", "pulses": "Pulses", "nuts": "Nuts", "seeds": "Seeds",
             "alternative_products": "Alternative products", "tofu_tempeh_seitan": "Tofu, tempeh & seitan",
             "dairy": "Dairy", "dairy_alternatives": "Dairy alternatives", "animal_fat": "Animal-based fat",
             "plant_fat": "Plant-based fat", "whole_grain": "Whole grain", "refined_grain": "Refined grain"}
FG_NAME = {"FG1": "Protein foods", "FG2": "Dairy and alternatives", "FG3": "Fats", "FG4": "Fruit and vegetables",
           "FG5": "Grains", "FG6": "Starchy vegetables", "FG7": "Discretionary"}
CAT_LABEL = {"meat_based": "meat-based", "seafood_based": "seafood-based", "vegetarian": "vegetarian", "vegan": "vegan"}
SPLIT_LABEL = {"one_side": "All ingredients on one side", "class": "Follows the class",
               "step4a": "First ingredient", "ingredients": "Ingredient list",
               "nevo": "NEVO food table (ratio only)", "convention": "50/50 convention"}
PROT_LABEL = {"declared": "Declared by the retailer", "reference_base": "Altera reference base",
              "off_ean": "Open Food Facts and NEVO, by EAN", "name_match": "National food table (CIQUAL in France)",
              "stratum_median": "Median of the family"}
PROCESSED = {"Cooked Meats", "Cured Meats", "Pâté", "Sausages & Burgers"}
CHEESE = {"Fresh Cheese", "Hard Cheese", "Soft Cheese"}
MILK = {"Milk", "Milk Drinks", "Yoghurt"}

TOP_FAMILIES = 40
TOP_PRODUCTS = 50


def detail(p: dict) -> str:
    """The GPA sub-category of a product, inside its Protein Tracker class."""
    c, w, fam = p["pt"], p.get("wwf") or {}, p["family"]
    fg, sub, struct = w.get("fg"), w.get("sub"), w.get("struct")
    if c == "animal_core":
        if not p.get("wwf_in"):
            return "Other"
        if fam in PROCESSED:
            return "Dishes and processed products"
        if fg == "FG1":
            return {"red_meat": "Red meat", "poultry": "Poultry", "fish_shellfish": "Fish and seafood",
                    "eggs": "Eggs"}.get(sub, "Other animals (e.g. snails)")
        if fam in CHEESE:
            return "Cheese"
        if fam in MILK:
            return "Milk, yoghurt and cream"
        return "Discretionary products" if fg == "FG7" else "Other"
    if c == "composite":
        if struct == "composite":
            return {"meat_based": "Meat-based dishes", "seafood_based": "Fish-based dishes",
                    "vegetarian": "Vegetarian dishes", "vegan": "Vegan dishes"}[w["cat"]]
        if fg == "FG7":
            return "Discretionary products"
        return "Cereals and bread" if fg == "FG5" else "Other"
    if c == "plant_non_core":
        if not p.get("wwf_in"):
            return "Outside the WWF scope"
        if fg == "FG5":
            return "Whole grains" if sub == "whole_grain" else "Refined grains"
        return {"FG4": "Fruit and vegetables", "FG6": "Starchy foods", "FG7": "Discretionary products"}.get(fg, "Other")
    return {"pulses": "Pulses", "nuts": "Nuts", "seeds": "Seeds", "alternative_products": "Plant-based meat alternatives",
            "tofu_tempeh_seitan": "Tofu, tempeh, seitan", "dairy_alternatives": "Plant-based dairy alternatives"}.get(sub, "Other")


# ------------------------------------------------------------ the line-level rows
# Data tab columns: (key, header, width, number format). Keys after "weight_src" are formulas.
DATA_COLS = [
    ("id", "Product ID", 11, None), ("code", "EAN 13", 16, "@"), ("label", "Till label", 32, None),
    ("name", "Product name", 34, None), ("brand", "Brand", 16, None), ("dept", "Department", 20, None),
    ("family", "Family", 20, None), ("in_scope", "In scope", 9, None), ("pt", "Protein Tracker class", 22, None),
    ("detail", "Class detail (GPA)", 26, None), ("level", "Data level", 8, None), ("split_src", "Split source", 24, None),
    ("prot", "Protein (g/100 g)", 10, "0.0"), ("prot_src", "Protein source", 36, None),
    ("plant", "Plant protein (g/100 g)", 10, "0.00"), ("animal", "Animal protein (g/100 g)", 10, "0.00"),
    ("wwf_in", "WWF scope", 9, None), ("struct", "WWF structure", 11, None), ("fg", "Food group", 8, None),
    ("sub", "WWF sub-group", 20, None), ("src", "WWF origin", 9, None), ("cat", "Composite category", 14, None),
    ("excl", "WWF exclusion (Appendix A)", 16, None), ("dairy", "Dairy factor", 8, "0"),
    ("weight", "Unit weight (kg)", 10, "0.000"), ("weight_src", "Weight source", 13, None),
    ("units", "Units sold", 11, INT), ("kg_in", "kg sold (typed)", 11, KG), ("kg", "kg sold", 12, KG),
    ("mass", "Mass in scope (kg)", 12, KG), ("prot_kg", "Protein (kg)", 11, KG),
    ("plant_kg", "Plant protein (kg)", 11, KG), ("animal_kg", "Animal protein (kg)", 11, KG),
    ("wwf_t", "WWF mass (t)", 10, "0.000"), ("wwf_eq", "WWF mass, dairy eq. (t)", 11, "0.000"),
]
DC = {k: get_column_letter(i + 1) for i, (k, *_) in enumerate(DATA_COLS)}
D0 = 5                      # first product row on Data
W0 = 31                     # first product row on Welcome


def flat(p: dict) -> dict:
    """One product as the Data tab shows it, plus the values its formulas will compute."""
    w = p.get("wwf") or {}
    on, wwf = bool(p.get("in_scope")), bool(p.get("wwf_in"))
    by_weight = p["weight_src"] == "variable"                 # sold by weight: volumes typed in kilos
    units = p.get("units", 0)
    kg = round(units * p["weight"], 3)
    r = {
        "id": p["id"], "code": p["code"], "label": p["label"], "name": p["name"], "brand": p["brand"],
        "dept": p["dept"], "family": p["family"], "in_scope": "yes" if on else "no",
        "pt": PT_LABEL[p["pt"]] if on else "", "detail": detail(p) if on else "", "level": p.get("level") or "",
        "split_src": SPLIT_LABEL.get(p.get("split_src", ""), "") if on else "",
        "prot": p.get("prot", 0) if on else None, "prot_src": PROT_LABEL.get(p.get("prot_src", ""), "") if on else "",
        "plant": p.get("plant", 0) if on else None, "animal": p.get("animal", 0) if on else None,
        "wwf_in": "yes" if wwf else "no", "struct": w.get("struct", "") if wwf else "",
        "fg": w.get("fg", "") if wwf else "", "sub": SUB_LABEL.get(w.get("sub") or "", "") if wwf else "",
        "src": w.get("src", "") if wwf and w.get("struct") == "whole" else "",
        "cat": CAT_LABEL.get(w.get("cat") or "", "") if wwf else "",
        "excl": (w.get("excl") or "").replace("_", " "), "dairy": p.get("dairy", 1),
        "weight": p["weight"], "weight_src": p["weight_src"],
        # the two inputs, as pre-filled on Welcome
        "in_units": None if by_weight else units, "in_kg": kg if by_weight else None,
    }
    r["kg"] = kg if by_weight else units * p["weight"]
    r["mass"] = r["kg"] if on else 0.0
    r["prot_kg"] = r["mass"] * (r["prot"] or 0) / 100
    r["plant_kg"] = r["mass"] * (r["plant"] or 0) / 100
    r["animal_kg"] = r["mass"] * (r["animal"] or 0) / 100
    r["wwf_t"] = r["mass"] / 1000 if wwf else 0.0
    r["wwf_eq"] = r["wwf_t"] * r["dairy"] if r["fg"] == "FG2" else r["wwf_t"]
    r["_class"] = p["pt"] if on else ""
    return r


class Book:
    """The workbook plus, for every formula written, the value Python expects of it."""

    def __init__(self, rows: list[dict]):
        self.rows = rows
        self.n = len(rows)
        self.wb = Workbook()
        self.expect: list[tuple[str, str, float]] = []

    def rng(self, key: str) -> str:
        c = DC[key]
        return f"Data!${c}${D0}:${c}${D0 + self.n - 1}"

    def sumifs(self, val: str, scale: float = 1, **crit: str) -> tuple[str, float]:
        """A SUMIFS formula over the Data tab and the value Python finds for it."""
        parts = ",".join(f'{self.rng(k)},"{v}"' for k, v in crit.items())
        f = f"SUMIFS({self.rng(val)},{parts})" if crit else f"SUM({self.rng(val)})"
        if scale != 1:
            f = f"{f}/{scale:g}"
        v = sum(r[val] for r in self.rows if all(r[k] == x for k, x in crit.items())) / scale
        return f, v

    def put(self, ws, cell: str, formula: str, value: float, fmt: str | None = None, font: Font = ARIAL):
        c = ws[cell]
        c.value = formula if formula.startswith("=") else "=" + formula
        c.font = font
        if fmt:
            c.number_format = fmt
        self.expect.append((ws.title, cell, value))
        return c


# ------------------------------------------------------------ layout helpers
def text(ws, cell: str, value, font: Font = ARIAL, merge: str | None = None, align: Alignment | None = None):
    c = ws[cell]
    c.value = value
    c.font = font
    if merge:
        ws.merge_cells(f"{cell}:{merge}")
    if align:
        c.alignment = align
    return c


def bar(ws, row: int, label: str, fill: str = GREEN, last: str = "G", font: Font = WHITE_B, height: float = 25.5):
    for col in range(1, ws[f"{last}1"].column + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
    text(ws, f"A{row}", "  " + label, font, merge=f"{last}{row}", align=Alignment(vertical="center"))
    ws.row_dimensions[row].height = height


def header(ws, row: int, labels: list[str | None], col: int = 1, height: float = 30):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, HEAD_BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 0 else "center")
    ws.row_dimensions[row].height = height


def widths(ws, spec: dict[str, float]):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def chart(title: str, kind: str = "stacked", legend: str = "b") -> BarChart:
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", kind, 100
    ch.title = title
    ch.style = 10
    ch.width, ch.height = 15, 7.5
    ch.legend.position = legend
    return ch


def colour(s: Series, rgb: str):
    s.graphicalProperties.solidFill = rgb
    s.graphicalProperties.line.solidFill = rgb


def shades(rgb: str, n: int) -> list[str]:
    """n tints of one colour, from the colour itself towards white."""
    base = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    out = []
    for k in range(n):
        t = 0.65 * k / max(1, n - 1)
        out.append("".join(f"{round(c + (255 - c) * t):02X}" for c in base))
    return out


def split_row(bk: Book, ws, row: int, crit: dict, weight_shares: bool):
    """Mass, protein, plant and animal protein (t), then the plant / animal shares."""
    fm, vm = bk.sumifs("mass", 1000, **crit)
    fp, vp = bk.sumifs("prot_kg", 1000, **crit)
    fpl, vpl = bk.sumifs("plant_kg", 1000, **crit)
    fa, va = bk.sumifs("animal_kg", 1000, **crit)
    for col, f, v in (("B", fm, vm), ("C", fp, vp), ("D", fpl, vpl), ("E", fa, va)):
        bk.put(ws, f"{col}{row}", f, v, T)
    if weight_shares is None:
        return
    if weight_shares:            # WWF reasons in weight: share of the mass that is of plant origin
        fpm, vpm = bk.sumifs("mass", 1000, **crit, src="plant")
        fam, vam = bk.sumifs("mass", 1000, **crit, src="animal")
        bk.put(ws, f"F{row}", f"IFERROR(({fpm})/B{row},0)", vpm / vm if vm else 0, PCT)
        bk.put(ws, f"G{row}", f"IFERROR(({fam})/B{row},0)", vam / vm if vm else 0, PCT)
    else:                        # Protein Tracker reasons in protein
        bk.put(ws, f"F{row}", f"IFERROR(D{row}/C{row},0)", vpl / vp if vp else 0, PCT)
        bk.put(ws, f"G{row}", f"IFERROR(E{row}/C{row},0)", va / vp if vp else 0, PCT)


def sum_rows(bk: Book, ws, cell: str, rows: list[int], col: str, fmt: str = T, font: Font = BOLD):
    refs = [f"{col}{r}" for r in rows]
    vals = {(t, c): v for t, c, v in bk.expect}
    total = sum(vals[(ws.title, r)] for r in refs)
    bk.put(ws, cell, "+".join(refs) if len(refs) < 6 else f"SUM({','.join(refs)})", total, fmt, font)
    return total


# ------------------------------------------------------------ Welcome
def welcome(bk: Book, order: list[dict], n_fam: int, n_dept: int, n_brand: int):
    ws = bk.wb.active
    ws.title = "Welcome"
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 24, "B": 18, "C": 46, "D": 16, "E": 16, "F": 3, "G": 16, "H": 18})
    logo = Image(str(ASSETS / "altera_wordmark.png"))
    logo.width, logo.height = 166, 67
    ws.add_image(logo, "A1")
    hm = Image(str(ASSETS / "harbour_market_logo.png"))
    hm.width, hm.height = 84, 67
    ws.add_image(hm, "B1")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    text(ws, "C1", "PROTEIN RATIO\nHarbour Market catalogue", Font(name="Arial", size=20, bold=True, color=NAVY),
         merge="F2", align=Alignment(wrap_text=True, vertical="center"))
    text(ws, "A3", BANNER, WARN, merge="H3", align=Alignment(vertical="center"))
    ws.row_dimensions[3].height = 22

    para = [
        (5, "What this workbook contains",
         "The share of the protein you put on the market that is of plant origin, and how it breaks down — "
         "by the WWF Planet-Based Diets methodology first, then by Protein Tracker."),
        (7, "How to use it",
         "Type your volumes in the yellow columns below: units sold, or kilos sold for products sold by weight. "
         "When both are filled, the kilos win. Every tab recalculates; there is nothing else to enter."),
        (9, "Volumes are pre-filled",
         "The volumes shown are our estimates. On a real engagement, when the sales file arrives without volumes, "
         "Altera deduces them from your sales rank in kilos; your own figures replace them. Here they are invented."),
    ]
    for row, k, v in para:
        text(ws, f"A{row}", k, NAVY_B, align=Alignment(wrap_text=True, vertical="top"))
        text(ws, f"B{row}", v, ARIAL, merge=f"F{row}", align=WRAP_TOP)
        ws.row_dimensions[row].height = 30

    text(ws, "A11", "THE TABS", NAVY_B)
    tabs = [
        ("WWF", "the seven food groups, whole against composite, the detail by sub-group"),
        ("Protein Tracker", "the ratio and the four classes of the GPA methodology"),
        ("By family", f"the {TOP_FAMILIES} families carrying the most protein, out of {n_fam}"),
        ("By department", f"the {n_dept} departments"),
        ("By brand", f"the {n_brand} brands"),
        ("Top and flop", f"the {TOP_PRODUCTS} most plant-based products and the {TOP_PRODUCTS} least"),
        ("Data", "one row per product: classes, contents, split, data level and the source of every value"),
    ]
    for i, (k, v) in enumerate(tabs, start=12):
        text(ws, f"A{i}", k, Font(name="Arial", size=10, bold=True, color=NAVY, underline="single"))
        ws[f"A{i}"].hyperlink = Hyperlink(ref=f"A{i}", location=f"'{k}'!A1", display=k)
        text(ws, f"B{i}", v, ARIAL, merge=f"F{i}")

    # accuracy: how it is measured; this invented catalogue has no figure to show
    text(ws, "A20", "ACCURACY", NAVY_B, align=Alignment(vertical="top"))
    text(ws, "B20", "Measured\nby hand", Font(name="Arial", size=18, bold=True, color=GREEN), merge="B26",
         align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
    text(ws, "C20",
         "On a real engagement, this box gives the accuracy measured on your own catalogue, one rate per field: "
         "whole or composite, food group, sub-group, Protein Tracker class.\n"
         "It is measured on a control sample, sealed before review:\n"
         "  •  500 rows drawn at random\n"
         "  •  the 500 products weighing most in tonnes of protein\n"
         "  •  the 3 best sellers of every category\n"
         "Every row of the sample is then reviewed by hand: no model grades a model.\n"
         "Harbour Market is invented, so this copy shows no accuracy figure.",
         ARIAL, merge="F26", align=WRAP_TOP)
    for r in range(20, 27):
        ws.row_dimensions[r].height = 17

    bar(ws, 28, "OUR ESTIMATED VOLUMES — type yours over them", fill=ORANGE, last="H",
        font=Font(name="Arial", size=12, bold=True, color="FFFFFF"))
    text(ws, "A29", "Sorted by our estimated rank, deduced from the volumes we estimated. Your own ranks, on the right, "
         "are recalculated from the yellow columns: where they differ from ours, our estimate departs.",
         NOTE, merge="H29", align=WRAP_TOP)
    ws.row_dimensions[29].height = 30
    header(ws, 30, ["Estimated rank (Altera)", "EAN 13", "Product", "Units sold", "or kg sold", None,
                    "Your sales rank (kilos)", "Your sales rank (quantity)"], height=39.75)
    ws["F30"].fill, ws["F30"].border = PatternFill(fill_type=None), Border()

    last = W0 + bk.n - 1
    kg_rng = bk.rng("kg")
    for i, r in enumerate(order):
        row = W0 + i
        ws[f"A{row}"] = i + 1
        ws[f"B{row}"] = r["code"]
        ws[f"B{row}"].number_format = "@"
        ws[f"C{row}"] = r["name"]
        for col, v, fmt in (("D", r["in_units"], INT), ("E", r["in_kg"], KG)):
            c = ws[f"{col}{row}"]
            c.value, c.fill, c.border, c.number_format = v, INPUT_FILL, INPUT_BORDER, fmt
        d = D0 + i
        bk.put(ws, f"G{row}", f'IF(Data!{DC["kg"]}{d}>0,RANK(Data!{DC["kg"]}{d},{kg_rng}),"")',
               sum(1 for x in order if x["kg"] > r["kg"]) + 1 if r["kg"] > 0 else "", INT)
        bk.put(ws, f"H{row}", f'IF(N(D{row})>0,RANK(D{row},$D${W0}:$D${last}),"")',
               sum(1 for x in order if (x["in_units"] or 0) > (r["in_units"] or 0)) + 1 if (r["in_units"] or 0) > 0 else "",
               INT)
        for col in "ABCDE":
            ws[f"{col}{row}"].font = ARIAL
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="A volume is a number, zero or more.", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"D{W0}:E{last}")


# ------------------------------------------------------------ Data
def data(bk: Book, order: list[dict]):
    ws = bk.wb.create_sheet("Data")
    ws.sheet_view.showGridLines = False
    bar(ws, 1, "DATA — one row per product", fill=NAVY, last=get_column_letter(len(DATA_COLS)))
    text(ws, "A2", BANNER + " Volumes come from the yellow columns of the Welcome tab; every other tab reads this one.",
         SMALL)
    header(ws, 4, [h for _, h, _, _ in DATA_COLS])
    for i, (k, _, w, _) in enumerate(DATA_COLS):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    for i, r in enumerate(order):
        d, wrow = D0 + i, W0 + i
        for ci, (k, _, _, fmt) in enumerate(DATA_COLS, start=1):
            if k in ("units", "kg_in", "kg", "mass", "prot_kg", "plant_kg", "animal_kg", "wwf_t", "wwf_eq"):
                continue
            v = r[k]
            c = ws.cell(row=d, column=ci, value=None if v == "" else v)
            c.font = ARIAL
            if fmt:
                c.number_format = fmt
        L = DC
        f = {
            "units": f'IF(ISBLANK(Welcome!D{wrow}),"",Welcome!D{wrow})',
            "kg_in": f'IF(ISBLANK(Welcome!E{wrow}),"",Welcome!E{wrow})',
            "kg": f"IF(N({L['kg_in']}{d})>0,N({L['kg_in']}{d}),N({L['units']}{d})*{L['weight']}{d})",
            "mass": f'IF({L["in_scope"]}{d}="yes",{L["kg"]}{d},0)',
            "prot_kg": f"{L['mass']}{d}*N({L['prot']}{d})/100",
            "plant_kg": f"{L['mass']}{d}*N({L['plant']}{d})/100",
            "animal_kg": f"{L['mass']}{d}*N({L['animal']}{d})/100",
            "wwf_t": f'IF({L["wwf_in"]}{d}="yes",{L["mass"]}{d}/1000,0)',
            "wwf_eq": f'IF({L["fg"]}{d}="FG2",{L["wwf_t"]}{d}*{L["dairy"]}{d},{L["wwf_t"]}{d})',
        }
        for k, fml in f.items():
            fmt = next(x[3] for x in DATA_COLS if x[0] == k)
            ws[f"{L[k]}{d}"] = "=" + fml
            ws[f"{L[k]}{d}"].font = ARIAL
            ws[f"{L[k]}{d}"].number_format = fmt
    ws.freeze_panes = f"D{D0}"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(DATA_COLS))}{D0 + bk.n - 1}"


# ------------------------------------------------------------ WWF
PHD = {21: 0.10, 27: 0.06, 34: 0.19, 35: "N/A", 37: 0.01, 38: 0.03, 40: 0.38, 42: 0.18, 43: 0.0, 45: 0.04,
       47: "N/A", 48: "N/A", 53: 0.60, 54: 0.40, 56: "N/A", 57: "N/A", 59: 0.74, 60: 0.26}


def wwf(bk: Book, n_cheese: int):
    ws = bk.wb.create_sheet("WWF")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 46, "B": 15, "C": 12, "D": 10, "E": 13, "F": 15, "G": 12, "H": 3, "I": 14, "J": 12})
    bar(ws, 1, "WWF — PLANET-BASED DIETS · APPENDIX D FORMAT REPORT", fill=NAVY, last="J")
    for row, k, v in ((3, "METHODOLOGY", "WWF Planet-Based Diets"),
                      (4, "INPUT DATA", f"Harbour Market catalogue (fictitious), {bk.n} products, volumes entered on the Welcome tab"),
                      (5, "STEP", "Step 1 — product level")):
        text(ws, f"A{row}", k, LABEL)
        text(ws, f"B{row}", v, ARIAL, merge=f"J{row}")

    text(ws, "A8", "Plant share of protein foods", NAVY_B, merge="B8")
    text(ws, "A9", "on WWF Food Group 1, in dairy equivalents", SMALL, merge="B9")
    text(ws, "A11", "Plant share of the whole diet", NAVY_B, merge="B11")
    text(ws, "A12", "on every whole product in the WWF scope", SMALL, merge="B12")

    bar(ws, 15, "WHOLE PRODUCTS", last="J")
    text(ws, "A17", "Three units, three readings. Dairy equivalents are the unit the WWF methodology imposes; protein is "
         "what this study is about. They do not rank food groups in the same order.", NOTE, merge="J17", align=WRAP_TOP)
    ws.row_dimensions[17].height = 30
    for a, b, lab in (("B", "D", "Dairy equivalents"), ("F", "G", "Product mass"), ("I", "J", "Protein")):
        for col in range(ws[f"{a}18"].column, ws[f"{b}18"].column + 1):
            ws.cell(row=18, column=col).fill = BAND_FILL
        text(ws, f"{a}18", f"{lab} — Harbour Market", Font(name="Arial", size=9, bold=True, color=NAVY),
             merge=f"{b}18", align=Alignment(horizontal="center"))
    header(ws, 19, [None, "Mass (t)", "Share", "PHD", None, "Mass (t)", "Share", None, "Protein (t)", "Share"])
    for col in "EH":
        ws[f"{col}19"].fill, ws[f"{col}19"].border = PatternFill(fill_type=None), Border()

    W = dict(struct="whole")
    # (row, label, criteria or None for a heading); criteria select whole products on the Data tab
    lines = [
        (20, "FOOD GROUP 1", None), (21, "Plant-Based", dict(fg="FG1", src="plant")),
        (22, "    Pulses", dict(fg="FG1", sub="Pulses")), (23, "    Nuts", dict(fg="FG1", sub="Nuts")),
        (24, "    Seeds", dict(fg="FG1", sub="Seeds")),
        (25, "    Alternative Products", dict(fg="FG1", sub="Alternative products")),
        (26, "    Tofu, tempeh & seitan", dict(fg="FG1", sub="Tofu, tempeh & seitan")),
        (27, "Animal-Based", dict(fg="FG1", src="animal")),
        (28, "    Red meat", dict(fg="FG1", sub="Red meat")), (29, "    Poultry", dict(fg="FG1", sub="Poultry")),
        (30, "    Fish & Shellfish", dict(fg="FG1", sub="Fish & shellfish")), (31, "    Eggs", dict(fg="FG1", sub="Eggs")),
        (32, "    Other (e.g. snails)", dict(fg="FG1", sub="Other animal")),
        (33, "FOOD GROUP 2", None), (34, "Dairy", dict(fg="FG2", src="animal")),
        (35, "Dairy Alternatives", dict(fg="FG2", src="plant")),
        (36, "FOOD GROUP 3", None), (37, "Animal-Based Fats", dict(fg="FG3", src="animal")),
        (38, "Plant-Based Fats", dict(fg="FG3", src="plant")),
        (39, "FOOD GROUP 4", None), (40, "Fruits & Vegetables", dict(fg="FG4")),
        (41, "FOOD GROUP 5", None), (42, "Whole Grains", dict(fg="FG5", sub="Whole grain")),
        (43, "Refined Grains", dict(fg="FG5", sub="Refined grain")),
        (44, "FOOD GROUP 6", None), (45, "Starchy Vegetables", dict(fg="FG6")),
        (46, "FOOD GROUP 7", None), (47, "Plant-Based", dict(fg="FG7", src="plant")),
        (48, "Animal-Based", dict(fg="FG7", src="animal")),
    ]
    units = (("B", "wwf_eq", 1), ("F", "wwf_t", 1), ("I", "prot_kg", 1000))
    tot = {}
    for col, val, sc in units:
        f, v = bk.sumifs(val, sc, **W)
        tot[col] = (f, v)
    got: dict[tuple[str, int], float] = {}
    for row, lab, crit in lines:
        if crit is None:
            for col in range(1, 11):
                ws.cell(row=row, column=col).fill = BAND_FILL
            text(ws, f"A{row}", lab, Font(name="Arial", size=10, bold=True, color=NAVY))
            continue
        text(ws, f"A{row}", lab, BOLD if not lab.startswith(" ") else ARIAL)
        for col, val, sc in units:
            f, v = bk.sumifs(val, sc, **W, **crit)
            got[(col, row)] = v
            bk.put(ws, f"{col}{row}", f, v, T)
            share = chr(ord(col) + 1)
            bk.put(ws, f"{share}{row}", f"IFERROR({col}{row}/{col}$51,0)", v / tot[col][1] if tot[col][1] else 0, PCT)
        ws[f"D{row}"] = PHD.get(row, "—")
        ws[f"D{row}"].number_format, ws[f"D{row}"].font = "0%", Font(name="Arial", size=10, color=GREY)
        ws[f"D{row}"].alignment = Alignment(horizontal="center")
    # what the classification could not settle: none here, but the rows stay so a real file can show them
    text(ws, "A49", "Undetermined origin", ARIAL)
    text(ws, "A50", "Undetermined sub-group", ARIAL)
    for col, val, sc in units:
        f_all, v_all = tot[col]
        fp, vp = bk.sumifs(val, sc, **W, src="plant")
        fa, va = bk.sumifs(val, sc, **W, src="animal")
        bk.put(ws, f"{col}49", f"{f_all}-{fp}-{fa}", v_all - vp - va, T)
        f1, v1 = bk.sumifs(val, sc, **W, fg="FG1")
        f5, v5 = bk.sumifs(val, sc, **W, fg="FG5")
        subs = sum(got[(col, r)] for r in (22, 23, 24, 25, 26, 28, 29, 30, 31, 32, 42, 43))
        bk.put(ws, f"{col}50", f"{f1}+{f5}-SUM({col}22:{col}26)-SUM({col}28:{col}32)-{col}42-{col}43", v1 + v5 - subs, T)
        share = chr(ord(col) + 1)
        for r in (49, 50):
            v = [x for t, c, x in bk.expect if t == "WWF" and c == f"{col}{r}"][0]
            bk.put(ws, f"{share}{r}", f"IFERROR({col}{r}/{col}$51,0)", v / v_all if v_all else 0, PCT)
        bk.put(ws, f"{col}51", f_all, v_all, T, BOLD)
        bk.put(ws, f"{share}51", f"IFERROR({col}51/{col}$51,0)", 1.0 if v_all else 0, PCT, BOLD)
    for r in (49, 50):
        ws[f"D{r}"] = "—"
        ws[f"D{r}"].font, ws[f"D{r}"].alignment = Font(name="Arial", size=10, color=GREY), Alignment(horizontal="center")
    text(ws, "A51", "Total — whole products", BOLD)

    # proportions by source, in dairy equivalents
    B = {r: got[("B", r)] for r in (21, 27, 34, 35, 37, 38, 40, 42, 43, 45, 47, 48)}
    plant_rows, animal_rows = (21, 35, 38, 40, 42, 43, 45, 47), (27, 34, 37, 48)
    vp, va = sum(B[r] for r in plant_rows), sum(B[r] for r in animal_rows)
    props = [
        (52, "FOOD GROUP 1 PROPORTION BY SOURCE", None, None),
        (53, "    Plant", "IFERROR(B21/(B21+B27),0)", B[21] / (B[21] + B[27])),
        (54, "    Animal", "IFERROR(B27/(B21+B27),0)", B[27] / (B[21] + B[27])),
        (55, "FOOD GROUP 2 PROPORTION BY SOURCE", None, None),
        (56, "    Plant", "IFERROR(B35/(B34+B35),0)", B[35] / (B[34] + B[35])),
        (57, "    Animal", "IFERROR(B34/(B34+B35),0)", B[34] / (B[34] + B[35])),
        (58, "WHOLE DIET PROPORTION BY SOURCE", None, None),
        (59, "    Plant", f"IFERROR(SUM({','.join(f'B{r}' for r in plant_rows)})/(SUM({','.join(f'B{r}' for r in plant_rows)})"
                        f"+SUM({','.join(f'B{r}' for r in animal_rows)})),0)", vp / (vp + va)),
        (60, "    Animal", "IFERROR(1-C59,0)", va / (vp + va)),
    ]
    for row, lab, f, v in props:
        if f is None:
            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = BAND_FILL
            text(ws, f"A{row}", lab, Font(name="Arial", size=10, bold=True, color=NAVY))
            continue
        text(ws, f"A{row}", lab, ARIAL)
        bk.put(ws, f"C{row}", f, v, PCT, Font(name="Arial", size=11))
        ws[f"D{row}"] = PHD[row]
        ws[f"D{row}"].number_format, ws[f"D{row}"].font = "0%", Font(name="Arial", size=10, color=GREY)
        ws[f"D{row}"].alignment = Alignment(horizontal="center")
    for cell, src, v in (("C8", "C53", B[21] / (B[21] + B[27])), ("C11", "C59", vp / (vp + va))):
        bk.put(ws, cell, src, v, PCT, KPI)
        r = int(cell[1:])
        ws.merge_cells(f"C{r}:D{r + 1}")
        ws[cell].alignment = Alignment(vertical="center")

    # composite products
    bar(ws, 62, "COMPOSITE PRODUCTS", last="J")
    fw, vw = bk.sumifs("wwf_t", 1, struct="whole")
    fc, vc = bk.sumifs("wwf_t", 1, struct="composite")
    text(ws, "A64", "Share of composites in measured sales", NAVY_B, merge="B64")
    bk.put(ws, "C64", f"IFERROR({fc}/({fc}+{fw}),0)", vc / (vc + vw), PCT, Font(name="Arial", size=14, bold=True, color=GREEN))
    header(ws, 66, ["COMPOSITION", "Mass (t)", "Share"])
    cats = [("Meat-based", "meat-based"), ("Seafood-based", "seafood-based"), ("Vegetarian", "vegetarian"), ("Vegan", "vegan")]
    for i, (lab, key) in enumerate(cats, start=67):
        text(ws, f"A{i}", lab)
        f, v = bk.sumifs("wwf_t", 1, struct="composite", cat=key)
        bk.put(ws, f"B{i}", f, v, T)
        bk.put(ws, f"C{i}", f"IFERROR(B{i}/B$71,0)", v / vc, PCT)
    text(ws, "A71", "Total", BOLD)
    bk.put(ws, "B71", "SUM(B67:B70)", vc, T, BOLD)
    bk.put(ws, "C71", "SUM(C67:C70)", 1.0, PCT, BOLD)

    notes = [
        (74, "PHD column: the Planetary Health Diet proportions as printed in the “PHD” column of Appendix D. "
             "A dash marks a row for which Appendix D gives no target."),
        (75, f"Appendix B dairy equivalents are applied: {n_cheese} cheeses at a factor of 1:10, everything else at 1:1. "
             "The product-mass column gives the mass actually sold, without the conversion; the gap between the two "
             "columns is the conversion itself."),
        (76, "Composites are counted at the full weight of the product, as the “composite product level” approach of "
             "step 1 prescribes."),
    ]
    for row, t in notes:
        text(ws, f"A{row}", t, SMALL, merge=f"G{row}", align=WRAP_TOP)
        ws.row_dimensions[row].height = 26

    bar(ws, 79, "THE SAME BREAKDOWN, IN PROTEIN", last="J")
    notes = [
        (81, "Appendix D reasons in tonnes of product. The table below takes the seven food groups again, in tonnes of protein."),
        (82, "Mass and protein columns are in PRODUCT mass: the dairy equivalents of Appendix B do not apply here. They "
             "turn a weight of cheese into a weight of milk, which means nothing for protein: it is the same on both sides."),
        (83, "Plant and animal shares are computed on WEIGHT, as the WWF methodology asks. The share of protein is read "
             "on the Protein Tracker tab, whose unit it is."),
    ]
    for row, t in notes:
        text(ws, f"A{row}", t, NOTE, merge=f"G{row}", align=WRAP_TOP)
        ws.row_dimensions[row].height = 26
    cols = ["Mass (t)", "Protein (t)", "Plant protein (t)", "Animal protein (t)", "Plant share (weight)", "Animal share (weight)"]
    header(ws, 85, ["Food group", *cols])
    for i, g in enumerate(FG_NAME, start=86):
        text(ws, f"A{i}", f"{g} · {FG_NAME[g]}")
        split_row(bk, ws, i, dict(struct="whole", fg=g), True)
    text(ws, "A93", "Total", BOLD)
    for col in "BCDE":
        sum_rows(bk, ws, f"{col}93", list(range(86, 93)), col)
    vals = {c: v for t, c, v in bk.expect if t == "WWF"}
    fpm, vpm = bk.sumifs("mass", 1000, struct="whole", src="plant")
    fam, vam = bk.sumifs("mass", 1000, struct="whole", src="animal")
    bk.put(ws, "F93", f"IFERROR(({fpm})/B93,0)", vpm / vals["B93"], PCT, BOLD)
    bk.put(ws, "G93", f"IFERROR(({fam})/B93,0)", vam / vals["B93"], PCT, BOLD)

    text(ws, "A95", "The rest of the catalogue", NAVY_B)
    header(ws, 96, [None, *cols])
    text(ws, "A97", "Composite products")
    split_row(bk, ws, 97, dict(struct="composite"), None)
    text(ws, "A98", "Outside the WWF scope")
    split_row(bk, ws, 98, dict(in_scope="yes", wwf_in="no"), None)

    ch = chart("Protein by food group")
    for col, rgb in (("D", GREEN), ("E", "C00000")):
        s = Series(Reference(ws, range_string=f"WWF!${col}$86:${col}$92"), title_from_data=False,
                   title=ws[f"{col}85"].value)
        colour(s, rgb)
        ch.series.append(s)
    ch.set_categories(Reference(ws, range_string="WWF!$A$86:$A$92"))
    ch.y_axis.title = "t of protein"
    ws.add_chart(ch, "A100")


# ------------------------------------------------------------ Protein Tracker
def protein_tracker(bk: Book):
    ws = bk.wb.create_sheet("Protein Tracker")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 42, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15})
    bar(ws, 1, "PROTEIN TRACKER — Green Protein Alliance", fill=NAVY)
    text(ws, "A3", "Plant share of protein", NAVY_B, merge="B3")
    text(ws, "A4", "on the Protein Tracker scope", SMALL, merge="B4")
    fpl, vpl = bk.sumifs("plant_kg", 1000)
    fp, vp = bk.sumifs("prot_kg", 1000)
    bk.put(ws, "C3", f"IFERROR(({fpl})/({fp}),0)", vpl / vp, PCT, KPI)
    ws.merge_cells("C3:D4")
    ws["C3"].alignment = Alignment(vertical="center")
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 24
    text(ws, "A6", "Every product falls into one class and one only. The four classes below cover the whole scope.",
         NOTE, merge="G6")

    bar(ws, 8, "THE FOUR CLASSES")
    header(ws, 10, [None, "Protein (t)", "Share"])
    for i, c in enumerate(PT_ORDER, start=11):
        text(ws, f"A{i}", PT_LABEL[c], BOLD)
        f, v = bk.sumifs("prot_kg", 1000, pt=PT_LABEL[c])
        bk.put(ws, f"B{i}", f, v, T)
        bk.put(ws, f"C{i}", f"IFERROR(B{i}/B$15,0)", v / vp, PCT)
    text(ws, "A15", "Total", BOLD)
    bk.put(ws, "B15", "SUM(B11:B14)", vp, T, BOLD)
    bk.put(ws, "C15", "SUM(C11:C14)", 1.0, PCT, BOLD)

    bar(ws, 18, "DETAIL OF EACH CLASS")
    text(ws, "A20", "The same total, broken down to the level the GPA methodology describes: meat, cheese, dairy, eggs, "
         "fish for the animal class, and likewise for the other three.", NOTE, merge="G20", align=WRAP_TOP)
    ws.row_dimensions[20].height = 28
    header(ws, 22, [None, "Protein (t)", "Share"])
    row = 23
    detail_rows: dict[str, list[int]] = {}
    for c in PT_ORDER:
        for col in range(1, 4):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=CLASS_FILL[c])
        text(ws, f"A{row}", PT_LABEL[c], Font(name="Arial", size=10, bold=True, color="FFFFFF"))
        row += 1
        detail_rows[c] = []
        for d in DETAIL[c]:
            text(ws, f"A{row}", "    " + d)
            f, v = bk.sumifs("prot_kg", 1000, pt=PT_LABEL[c], detail=d)
            bk.put(ws, f"B{row}", f, v, T)
            bk.put(ws, f"C{row}", f"IFERROR(B{row}/B$15,0)", v / vp, PCT)
            detail_rows[c].append(row)
            row += 1
    total_row = row
    text(ws, f"A{total_row}", "Total", BOLD)
    all_rows = [r for c in PT_ORDER for r in detail_rows[c]]
    sum_rows(bk, ws, f"B{total_row}", all_rows, "B")
    bk.put(ws, f"C{total_row}", f"IFERROR(B{total_row}/B$15,0)", 1.0, PCT, BOLD)

    sec = total_row + 3
    bar(ws, sec, "MASS AND PLANT / ANIMAL SPLIT")
    cols = ["Mass (t)", "Protein (t)", "Plant protein (t)", "Animal protein (t)", "Plant share", "Animal share"]
    header(ws, sec + 2, ["Class", *cols])
    r0 = sec + 3
    for i, c in enumerate(["plant_core", "plant_non_core", "composite", "animal_core"]):
        text(ws, f"A{r0 + i}", PT_SPLIT_LABEL[c])
        split_row(bk, ws, r0 + i, dict(pt=PT_LABEL[c]), False)
    tr = r0 + 4
    text(ws, f"A{tr}", "Total", BOLD)
    for col in "BCDE":
        sum_rows(bk, ws, f"{col}{tr}", list(range(r0, tr)), col)
    bk.put(ws, f"F{tr}", f"IFERROR(D{tr}/C{tr},0)", vpl / vp, PCT, BOLD)
    va = bk.sumifs("animal_kg", 1000)[1]
    bk.put(ws, f"G{tr}", f"IFERROR(E{tr}/C{tr},0)", va / vp, PCT, BOLD)
    text(ws, f"A{tr + 2}", "Out of scope", NAVY_B)
    header(ws, tr + 3, [None, *cols])
    text(ws, f"A{tr + 4}", "Out of scope — non-food and test items")
    out_kg = sum(r["kg"] for r in bk.rows if r["in_scope"] == "no")
    bk.put(ws, f"B{tr + 4}", f'SUMIFS({bk.rng("kg")},{bk.rng("in_scope")},"no")/1000', out_kg / 1000, T)
    for col in "CDEFG":
        ws[f"{col}{tr + 4}"] = "—"
        ws[f"{col}{tr + 4}"].font, ws[f"{col}{tr + 4}"].alignment = Font(name="Arial", size=10, color=GREY), Alignment(horizontal="center")

    ch = chart("Protein by class", "percentStacked", "r")
    for i, c in enumerate(PT_ORDER, start=11):
        s = Series(Reference(ws, range_string=f"'Protein Tracker'!$B${i}"), title=PT_LABEL[c])
        colour(s, CLASS_FILL[c])
        ch.series.append(s)
    ch.x_axis.delete = True
    ws.add_chart(ch, "H3")
    ch = chart("Protein by class, in detail", "percentStacked", "r")
    ch.width, ch.height = 20, 16
    for c in PT_ORDER:
        for r, rgb in zip(detail_rows[c], shades(CLASS_FILL[c], len(detail_rows[c]))):
            s = Series(Reference(ws, range_string=f"'Protein Tracker'!$B${r}"),
                       title=f"{PT_LABEL[c]} · {ws[f'A{r}'].value.strip()}")
            colour(s, rgb)
            ch.series.append(s)
    ch.x_axis.delete = True
    ws.add_chart(ch, "H22")


# ------------------------------------------------------------ By family, by department, by brand
GRID_COLS = ["Mass (t)", "Protein (t)", "Plant protein (t)", "Animal protein (t)", "Plant share", "Animal share"]


def ranked(bk: Book, key: str) -> list[tuple[str, float]]:
    agg = defaultdict(float)
    for r in bk.rows:
        if r["in_scope"] == "yes":
            agg[r[key]] += r["prot_kg"]
    return sorted(agg.items(), key=lambda kv: (-kv[1], kv[0]))


def grid(bk: Book, ws, start: int, key: str, first_label: str, top: int | None, noun: str):
    """One row per value of `key`, sorted by protein tonnage; the tail on one row; a total."""
    items = ranked(bk, key)
    shown = items if top is None else items[:top]
    header(ws, start, [first_label, *GRID_COLS])
    row = start + 1
    for name, _ in shown:
        text(ws, f"A{row}", name)
        split_row(bk, ws, row, {key: name, "in_scope": "yes"}, False)
        row += 1
    vals = lambda: {c: v for t, c, v in bk.expect if t == ws.title}  # noqa: E731
    fm, vm = bk.sumifs("mass", 1000)
    fp, vp = bk.sumifs("prot_kg", 1000)
    fpl, vpl = bk.sumifs("plant_kg", 1000)
    fa, va = bk.sumifs("animal_kg", 1000)
    rest = len(items) - len(shown)
    if rest:
        text(ws, f"A{row}", f"All the others ({rest} {noun})", Font(name="Arial", size=10, italic=True))
        v = vals()
        for col, f, tot in (("B", fm, vm), ("C", fp, vp), ("D", fpl, vpl), ("E", fa, va)):
            s = sum(v[f"{col}{r}"] for r in range(start + 1, row))
            bk.put(ws, f"{col}{row}", f"{f}-SUM({col}{start + 1}:{col}{row - 1})", tot - s, T)
        v = vals()
        bk.put(ws, f"F{row}", f"IFERROR(D{row}/C{row},0)", v[f"D{row}"] / v[f"C{row}"] if v[f"C{row}"] else 0, PCT)
        bk.put(ws, f"G{row}", f"IFERROR(E{row}/C{row},0)", v[f"E{row}"] / v[f"C{row}"] if v[f"C{row}"] else 0, PCT)
        row += 1
    text(ws, f"A{row}", "Total", BOLD)
    for col, tot in (("B", vm), ("C", vp), ("D", vpl), ("E", va)):
        bk.put(ws, f"{col}{row}", f"SUM({col}{start + 1}:{col}{row - 1})", tot, T, BOLD)
    bk.put(ws, f"F{row}", f"IFERROR(D{row}/C{row},0)", vpl / vp, PCT, BOLD)
    bk.put(ws, f"G{row}", f"IFERROR(E{row}/C{row},0)", va / vp, PCT, BOLD)
    for r in range(start + 1, row + 1):
        ws[f"A{r}"].alignment = Alignment(wrap_text=False)
    return len(items), row


def top15_chart(ws, title: str, head: int, anchor: str, n: int = 15):
    ch = chart(title)
    q = f"'{ws.title}'"
    for col, rgb in (("D", GREEN), ("E", "C00000")):
        s = Series(Reference(ws, range_string=f"{q}!${col}${head + 1}:${col}${head + n}"), title=ws[f"{col}{head}"].value)
        colour(s, rgb)
        ch.series.append(s)
    ch.set_categories(Reference(ws, range_string=f"{q}!$A${head + 1}:$A${head + n}"))
    ch.y_axis.title = "t of protein"
    ws.add_chart(ch, anchor)


def by_family(bk: Book) -> int:
    ws = bk.wb.create_sheet("By family")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 34, "B": 13, "C": 13, "D": 15, "E": 15, "F": 13, "G": 13})
    bar(ws, 1, "BY FAMILY", fill=NAVY)
    n_fam = len(ranked(bk, "family"))
    text(ws, "A3", f"Sorted by protein tonnage, from the pre-filled volumes. {TOP_FAMILIES} families shown out of {n_fam}; "
         f"the other {n_fam - TOP_FAMILIES} are grouped on one row.", NOTE, merge="G3")
    bar(ws, 5, "FIRST, THE COMPARABLE GROUPING")
    header(ws, 7, ["WWF sub-group", "Mass (t)", "Protein (t)", "Plant protein (t)", "Animal protein (t)",
                   "Plant share (weight)", "Animal share (weight)"])
    subs = [("FG1", "Red meat"), ("FG1", "Poultry"), ("FG1", "Fish & shellfish"), ("FG1", "Eggs"), ("FG1", "Other animal"),
            ("FG1", "Pulses"), ("FG1", "Nuts"), ("FG1", "Seeds"), ("FG1", "Tofu, tempeh & seitan"),
            ("FG1", "Alternative products"), ("FG2", "Dairy"), ("FG2", "Dairy alternatives"),
            ("FG3", "Animal-based fat"), ("FG3", "Plant-based fat"), ("FG5", "Whole grain"), ("FG5", "Refined grain")]
    for i, (g, s) in enumerate(subs, start=8):
        text(ws, f"A{i}", f"{g} · {s}")
        split_row(bk, ws, i, dict(struct="whole", fg=g, sub=s), True)
    text(ws, "A25", "The rest", NAVY_B)
    header(ws, 26, [None, *GRID_COLS[:4], "Plant share (weight)", "Animal share (weight)"])
    text(ws, "A27", "Composite products")
    split_row(bk, ws, 27, dict(struct="composite"), None)
    text(ws, "A28", "Outside the WWF scope")
    split_row(bk, ws, 28, dict(in_scope="yes", wwf_in="no"), None)
    text(ws, "A29", "Food groups FG4, FG6 and FG7 have no sub-group, so they do not appear above; the WWF tab details them.",
         SMALL, merge="G29")
    bar(ws, 32, "THEN, THE RETAILER'S GRID")
    grid(bk, ws, 34, "family", "Family", TOP_FAMILIES, "families")
    top15_chart(ws, "The 15 families carrying the most protein", 34, "I34")
    return n_fam


def by_simple(bk: Book, title: str, key: str, first: str, noun: str) -> int:
    ws = bk.wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 30, "B": 13, "C": 13, "D": 15, "E": 15, "F": 13, "G": 13})
    bar(ws, 1, title.upper(), fill=NAVY)
    n = len(ranked(bk, key))
    text(ws, "A3", f"Sorted by protein tonnage, from the pre-filled volumes. All {n} {noun} shown.", NOTE, merge="G3")
    grid(bk, ws, 5, key, first, None, noun)
    ws.freeze_panes = "A6"
    top15_chart(ws, f"The {min(15, n)} {noun} carrying the most protein", 5, "I6", min(15, n))
    return n


# ------------------------------------------------------------ Top and flop
def top_flop(bk: Book, order: list[dict]):
    ws = bk.wb.create_sheet("Top and flop")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 15, "B": 44, "C": 16, "D": 22, "E": 10, "F": 12, "G": 12, "H": 12, "I": 13, "J": 13, "K": 13})
    bar(ws, 1, "TOP AND FLOP — Protein Tracker ranking", fill=NAVY, last="K")
    text(ws, "A3", "Ranking basis: PROTEIN TRACKER. A product's plant share is its plant protein over its total protein; "
         "the WWF ranking, which reasons in weight, does not apply to a single product.", NOTE, merge="K3", align=WRAP_TOP)
    text(ws, "A4", "Weight, volume and protein per 100 g are read from the Data tab: these rows recalculate with your volumes.",
         NOTE, merge="K4")
    ws.row_dimensions[3].height = 28
    idx = {id(r): i for i, r in enumerate(order)}
    sold = [r for r in order if r["in_scope"] == "yes" and r["prot_kg"] > 0]
    share = lambda r: r["plant_kg"] / r["prot_kg"]  # noqa: E731
    top = sorted(sold, key=lambda r: (-share(r), -r["plant_kg"], r["id"]))[:TOP_PRODUCTS]
    flop = sorted(sold, key=lambda r: (share(r), -r["animal_kg"], r["id"]))[:TOP_PRODUCTS]
    heads = ["EAN 13", "Product", "Brand", "Family", "Food group", "Unit weight (kg)", "Units sold",
             "Protein (g/100 g)", "Plant protein (t)", "Animal protein (t)", "Plant share (protein)"]
    for sec, lab, items in ((6, f"The {TOP_PRODUCTS} most plant-based", top), (6 + TOP_PRODUCTS + 5, f"The {TOP_PRODUCTS} most animal", flop)):
        bar(ws, sec, lab, last="K", fill=GREEN if items is top else "C00000")
        header(ws, sec + 2, heads)
        for k, r in enumerate(items):
            row, d = sec + 3 + k, D0 + idx[id(r)]
            for col, v in zip("ABCDE", (r["code"], r["name"], r["brand"], r["family"], r["fg"] or "—")):
                text(ws, f"{col}{row}", v)
            ws[f"A{row}"].number_format = "@"
            ref = lambda key: f"Data!{DC[key]}{d}"  # noqa: E731
            bk.put(ws, f"F{row}", ref("weight"), r["weight"], "0.000")
            bk.put(ws, f"G{row}", ref("units"), r["in_units"] if r["in_units"] is not None else "", INT)
            bk.put(ws, f"H{row}", ref("prot"), r["prot"], "0.0")
            bk.put(ws, f"I{row}", f"{ref('plant_kg')}/1000", r["plant_kg"] / 1000, "#,##0.00")
            bk.put(ws, f"J{row}", f"{ref('animal_kg')}/1000", r["animal_kg"] / 1000, "#,##0.00")
            bk.put(ws, f"K{row}", f"IFERROR(I{row}/(I{row}+J{row}),0)", share(r), PCT)


# ------------------------------------------------------------ build, write, check
def build(products: list[dict]) -> Book:
    rows = [flat(p) for p in products]
    order = sorted(rows, key=lambda r: (-r["kg"], r["id"]))       # our estimated rank, heaviest first
    bk = Book(order)
    n_cheese = sum(1 for r in order if r["fg"] == "FG2" and r["dairy"] == 10)
    for r in order:
        if r["in_scope"] == "yes":
            assert r["detail"] in DETAIL[r["_class"]], (r["id"], r["detail"])
    # Welcome first so it opens first; it needs the counts the other tabs compute
    n_dept = len(ranked(bk, "dept"))
    n_brand = len(ranked(bk, "brand"))
    n_fam = len(ranked(bk, "family"))
    welcome(bk, order, n_fam, n_dept, n_brand)
    wwf(bk, n_cheese)
    protein_tracker(bk)
    by_family(bk)
    by_simple(bk, "By department", "dept", "Department", "departments")
    by_simple(bk, "By brand", "brand", "Brand", "brands")
    top_flop(bk, order)
    data(bk, order)
    bk.wb.active = 0
    bk.wb.save(XLSX)
    return bk


def write_csv(bk: Book) -> None:
    keys = [(h, k) for k, h, _, _ in DATA_COLS if k not in ("kg_in", "mass", "wwf_t", "wwf_eq")]
    with CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow([h for h, _ in keys])
        for r in sorted(bk.rows, key=lambda r: r["id"]):
            out = []
            for _, k in keys:
                v = r["in_units"] if k == "units" else r.get(k)
                out.append("" if v is None else (round(v, 4) if isinstance(v, float) else v))
            w.writerow(out)


def verify(bk: Book) -> list[str]:
    soffice = shutil.which("soffice")
    if not soffice:
        print("  LibreOffice not found: formulas NOT recalculated.")
        return []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(XLSX)],
                       check=True, capture_output=True, timeout=300)
        wb = load_workbook(Path(tmp) / XLSX.name, data_only=True)
    errors = []
    for sheet, cell, want in bk.expect:
        got = wb[sheet][cell].value
        if want == "" or want is None:
            ok = got in ("", None)
        else:
            ok = isinstance(got, (int, float)) and abs(got - want) <= 1e-6 * max(1, abs(want))
        if not ok:
            errors.append(f"{sheet}!{cell}: workbook {got!r} vs python {want!r}")
    bad = [f"{s.title}!{c.coordinate}" for s in wb.worksheets for row in s.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("#")]
    if bad:
        errors.append(f"{len(bad)} cells evaluate to an error, e.g. {bad[:5]}")
    print(f"  recalculated with LibreOffice: {len(bk.expect)} formula cells compared with Python")
    return errors


def main() -> int:
    products = json.loads(SRC.read_text())["products"]
    OUT.mkdir(parents=True, exist_ok=True)
    bk = build(products)
    write_csv(bk)
    errors = verify(bk)
    v = {(t, c): x for t, c, x in bk.expect}
    print(f"  WWF plant share FG1 {v[('WWF', 'C8')]:.1%} · whole diet {v[('WWF', 'C11')]:.1%} · "
          f"Protein Tracker {v[('Protein Tracker', 'C3')]:.1%}")
    print(f"  -> {XLSX.relative_to(ROOT)} ({XLSX.stat().st_size // 1024} KB), {CSV.relative_to(ROOT)} ({CSV.stat().st_size // 1024} KB)")
    if errors:
        print("REFUSED:", *errors[:30], sep="\n  ")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
